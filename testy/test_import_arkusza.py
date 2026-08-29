"""Strazniki importu z arkusza kalkulacyjnego.

Komplet regul sprawdzany jest na arkuszu SYNTETYCZNYM budowanym w tescie -
prawdziwy plik lezy w iCloud, poza repozytorium, i test na nim nie moglby byc
bramka na innej maszynie. Test na prawdziwym pliku istnieje osobno, pod
markerem `arkusz`, jako kontrola dwustronna uruchamiana swiadomie.

Kazdy arkusz syntetyczny odtwarza pulapki zmierzone w prawdziwym pliku:
naglowek w wierszu 11, dane od kolumny C, pusty ogon tabeli z zywymi
formulami ponizej danych, kwoty raz jako int raz jako float, sloty
-placeholdery poza zakresami tabel.
"""

import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from silnik.rdzen import budzet, import_arkusza, magazyn

SCIEZKA_PRAWDZIWEGO = Path.home() / (
    "Library/Mobile Documents/com~apple~CloudDocs/Finanse/"
    "Budżet Osobisty 2025 — kopia.xlsx")

#: Siedem wpisow: 3 typy, kwoty int i float, dwa CELOWO identyczne wiersze.
WPISY_SYNTETYCZNE = [
    (datetime.datetime(2025, 10, 1), "Wydatki", "Transport", 6),
    (datetime.datetime(2025, 10, 1), "Wydatki", "Zakupy", 13.99),
    (datetime.datetime(2025, 10, 1), "Oszczędności", "Reszta Revolut", 2.02),
    (datetime.datetime(2025, 10, 2), "Przychody", "Wynagrodzenie", 8130.78),
    (datetime.datetime(2025, 10, 3), "Wydatki", "Jedzenie poza domem", 12.00),
    (datetime.datetime(2025, 10, 3), "Wydatki", "Jedzenie poza domem", 12.00),
    (datetime.datetime(2025, 11, 1), "Wydatki", "Kredyty", 2400.33),
]


def _zbuduj_arkusz(sciezka: Path, wpisy=None) -> Path:
    """Buduje skoroszyt o ukladzie prawdziwego pliku, z jego pulapkami."""
    wpisy = WPISY_SYNTETYCZNE if wpisy is None else wpisy
    skoroszyt = openpyxl.Workbook()

    sledzenie = skoroszyt.active
    sledzenie.title = "Śledzenie Budżetu"
    sledzenie["C1"] = "Śledzenie Budżetu"
    for kolumna, naglowek in zip("CDEFGHI", [
            "Data", "Typ", "Kategoria", "Kwota", "Szczegóły", "Saldo",
            "Data Efektywna"]):
        sledzenie[f"{kolumna}11"] = naglowek
    for przesuniecie, (data, typ, kategoria, kwota) in enumerate(wpisy):
        wiersz = 12 + przesuniecie
        sledzenie[f"C{wiersz}"] = data
        sledzenie[f"D{wiersz}"] = typ
        sledzenie[f"E{wiersz}"] = kategoria
        sledzenie[f"F{wiersz}"] = kwota
        sledzenie[f"H{wiersz}"] = -22.01          # saldo - formula tablicowa
        sledzenie[f"I{wiersz}"] = data            # data efektywna - formula

    # PULAPKA: pusty ogon tabeli. Kolumny H i I maja tam zywe wartosci, wiec
    # `max_row` pokazuje wiersz 40 przy siedmiu realnych wpisach.
    for wiersz in range(12 + len(wpisy), 41):
        sledzenie[f"H{wiersz}"] = 0
        sledzenie[f"I{wiersz}"] = datetime.time(0, 0)
    sledzenie[f"D{12 + len(wpisy)}"] = "Wydatki"   # osierocony typ bez daty

    # PULAPKA DRUGA, wazniejsza: wiersz-sierota z PRAWDZIWA data pod pusta
    # przerwa. Bez niego straznik warunku stopu bylby atrapa - import
    # pomijajacy wiersze bez daty dawal by ten sam wynik co import
    # zatrzymujacy sie na pierwszym z nich, bo pusty ogon i tak nie ma dat.
    # Zatrzymanie czyta 7, skanowanie calosci czyta 8.
    sledzenie["C39"] = datetime.datetime(2030, 1, 1)
    sledzenie["D39"] = "Wydatki"
    sledzenie["E39"] = "Zakupy"
    sledzenie["F39"] = 99999

    planowanie = skoroszyt.create_sheet("Planowanie Budżetu")
    sekcje = {
        "Przychody": (9, ["Wynagrodzenie", "Inne"]),
        "Wydatki": (22, ["Transport", "Zakupy", "Jedzenie poza domem",
                         "Kredyty", "Inne"]),
        "Oszczędności": (41, ["Reszta Revolut", "Inwestycje"]),
    }
    for typ, (wiersz_naglowka, kategorie) in sekcje.items():
        planowanie[f"C{wiersz_naglowka}"] = typ
        planowanie[f"E{wiersz_naglowka}"] = datetime.datetime(2025, 10, 1)
        planowanie[f"F{wiersz_naglowka}"] = datetime.datetime(2025, 11, 1)
        for przesuniecie, kategoria in enumerate(kategorie):
            wiersz = wiersz_naglowka + 1 + przesuniecie
            planowanie[f"C{wiersz}"] = kategoria
            planowanie[f"E{wiersz}"] = 100 + przesuniecie
            planowanie[f"F{wiersz}"] = 200 + przesuniecie
        ostatni = wiersz_naglowka + len(kategorie)
        planowanie.add_table(Table(
            displayName=typ.replace("ś", "s").replace("ó", "o"),
            ref=f"C{wiersz_naglowka}:C{ostatni}"))
        # PULAPKA: slot-placeholder tuz pod zakresem tabeli. W prawdziwym pliku
        # jest ich osiem, w tym siedem powtorzen "Inne".
        planowanie[f"C{ostatni + 1}"] = "Wprowadź Kategorię…"
        planowanie[f"E{ostatni + 1}"] = 999

    skoroszyt.save(sciezka)
    return sciezka


@pytest.fixture
def arkusz(tmp_path):
    return _zbuduj_arkusz(tmp_path / "syntetyk.xlsx")


@pytest.fixture
def p(tmp_path):
    return magazyn.zmigruj_i_polacz(tmp_path / "import.db")


# --- kontrola dwustronna --------------------------------------------------

def test_syntetyk_naprawde_ma_siedem_wpisow(arkusz):
    """Bez tego "zaimportowano 7 z 7" bylo by zielone takze wtedy, gdy budowniczy
    arkusza przestanie cokolwiek zapisywac. Zero jest wtedy zgodne z zerem."""
    skoroszyt = openpyxl.load_workbook(arkusz, data_only=True)
    ws = skoroszyt["Śledzenie Budżetu"]

    ciagle = 0
    while isinstance(ws.cell(12 + ciagle, 3).value, datetime.datetime):
        ciagle += 1
    assert ciagle == 7

    wszystkie_daty = [w for w in range(12, ws.max_row + 1)
                      if isinstance(ws.cell(w, 3).value, datetime.datetime)]
    assert ws.max_row > 7 + 11, "pulapka pustego ogona ma byc w arkuszu"
    assert len(wszystkie_daty) == 8, (
        "pulapka wiersza-sieroty ma byc w arkuszu - bez niej straznik warunku "
        "stopu nie odroznia zatrzymania od pomijania")


# --- odczyt wpisow --------------------------------------------------------

def test_import_czyta_wszystkie_wpisy(arkusz, p):
    raport = import_arkusza.importuj(p, arkusz)
    assert raport.przeczytane == 7
    assert raport.zapisane == 7
    assert raport.odrzucone == 0
    assert len(magazyn.wczytaj_wpisy(p)) == 7


def test_import_zatrzymuje_sie_na_typie_daty_nie_na_max_row(arkusz, p):
    """`max_row` klamie: liczy pusty ogon z zywymi formulami w kolumnach H i I.

    Import oparty na tej liczbie wciagnalby dziesiatki pustych wierszy
    i wygladalby na udany.
    """
    skoroszyt = openpyxl.load_workbook(arkusz, data_only=True)
    assert skoroszyt["Śledzenie Budżetu"].max_row >= 40
    raport = import_arkusza.importuj(p, arkusz)
    assert raport.przeczytane == 7, (
        "import ma sie ZATRZYMAC na pierwszym wierszu bez daty, nie skanowac "
        "arkusza do konca - inaczej wciagnie wiersz-sierote z 2030 roku")
    assert Decimal("99999") not in set(magazyn.wczytaj_wpisy(p)["kwota"])


def test_import_czyta_od_kolumny_C_nie_A(arkusz, p):
    """Tabela w arkuszu jest wcieta o dwie kolumny. Przesuniecie o jedna daje
    date w miejscu typu i wszystkie wiersze wypadaja jako nieczytelne."""
    import_arkusza.importuj(p, arkusz)
    wpisy = magazyn.wczytaj_wpisy(p)
    assert set(wpisy["typ"]) <= set(budzet.TYPY)
    assert "Transport" in set(wpisy["kategoria"])


def test_kwoty_wchodza_jako_decimal_bez_straty_grosza(arkusz, p):
    """W prawdziwym pliku 58 kwot to `int`, a 146 to `float`. Droga przez
    `Decimal(float)` daje 13.990000000000000213... zamiast 13.99."""
    import_arkusza.importuj(p, arkusz)
    wpisy = magazyn.wczytaj_wpisy(p)
    kwoty = set(wpisy["kwota"])
    assert all(isinstance(k, Decimal) for k in kwoty)
    assert Decimal("13.99") in kwoty
    assert Decimal("6.00") in kwoty or Decimal("6") in kwoty
    assert Decimal("2400.33") in kwoty


def test_wartosc_logiczna_nie_jest_kwota():
    """`bool` jest w Pythonie podklasa `int`, wiec bez jawnego odrzucenia
    komorka z wartoscia PRAWDA weszlaby do bazy jako kwota 1,00 zl.

    Druga asercja to kontrola dwustronna: potwierdza, ze mechanizm w ogole
    cokolwiek przepuszcza. Sprawdzenie, do ktorego nic nie dociera, melduje
    sukces - a "None dla wszystkiego" spelnialoby pierwsza asercje idealnie.
    """
    assert import_arkusza._na_decimal(True) is None
    assert import_arkusza._na_decimal(False) is None
    assert import_arkusza._na_decimal(1) == Decimal("1.00")
    assert import_arkusza._na_decimal(13.99) == Decimal("13.99")

    # Komorka z TEKSTEM odrzucona - i to jest nosna polowa sprawdzenia typu.
    # `Decimal("13.99")` przeszlo by bez niej jako kwota, bo napis z kropka
    # jest dla Decimala poprawny; wyjatek lapie tylko zapis z przecinkiem.
    assert import_arkusza._na_decimal("13.99") is None
    assert import_arkusza._na_decimal("13,99") is None


def test_kwota_zaokraglana_z_zapisu_dziesietnego_nie_z_binarnego():
    """Jedyny przypadek, ktory ODROZNIA `Decimal(str(v))` od `Decimal(v)`.

    Przy dwoch miejscach po przecinku `quantize` zaciera roznice i oba zapisy
    daja to samo - dlatego pierwsza wersja tego testu przepuszczala mutacje.
    Przy trzech miejscach roznica jest widoczna: float trzyma 2,675 jako
    2,674999999999999822..., wiec droga przez liczbe binarna daje 2,67,
    a droga przez tekst - poprawne 2,68.

    W prawdziwym arkuszu wszystkie 204 kwoty maja dwa miejsca po przecinku,
    wiec dzis ta roznica nie wystepuje. Straznik pilnuje reguly na przyszlosc,
    nie biezacego pliku - i musi miec przypadek, ktory ja rusza.
    """
    assert import_arkusza._na_decimal(2.675) == Decimal("2.68")
    assert import_arkusza._na_decimal(2.665) == Decimal("2.67")


# --- raport ---------------------------------------------------------------

def test_import_melduje_liczby_a_nie_slowo_gotowe(arkusz, p):
    raport = import_arkusza.importuj(p, arkusz)
    assert "7" in raport.komunikat
    assert "gotowe" not in raport.komunikat.lower()
    assert isinstance(raport.powody, dict)


def test_import_podaje_powod_odrzucenia(tmp_path, p):
    zle = list(WPISY_SYNTETYCZNE) + [
        (datetime.datetime(2025, 10, 9), "Inwestycje", "Transport", 10),
        (datetime.datetime(2025, 10, 9), "Wydatki", "Czegoś takiego nie ma", 10),
    ]
    arkusz = _zbuduj_arkusz(tmp_path / "zle.xlsx", zle)
    raport = import_arkusza.importuj(p, arkusz)
    assert raport.przeczytane == 9
    assert raport.zapisane == 7
    assert raport.odrzucone == 2
    assert sum(raport.powody.values()) == 2
    assert any("Inwestycje" in powod for powod in raport.powody)


# --- deduplikacja wielozbiorowa -------------------------------------------

def test_dwa_identyczne_wpisy_w_pliku_wchodza_oba(arkusz, p):
    """Kolumna Szczegoly jest pusta we WSZYSTKICH 204 wierszach prawdziwego
    pliku, wiec klucz nie odroznia dwoch takich samych kaw tego samego dnia.

    Wiez UNIQUE zjadlby tu prawdziwe dane.
    """
    import_arkusza.importuj(p, arkusz)
    wpisy = magazyn.wczytaj_wpisy(p)
    jedzenie = wpisy[wpisy["kategoria"] == "Jedzenie poza domem"]
    assert len(jedzenie) == 2


def test_ponowny_import_niezmienionego_arkusza_zapisuje_zero(arkusz, p):
    pierwszy = import_arkusza.importuj(p, arkusz)
    drugi = import_arkusza.importuj(p, arkusz)
    assert pierwszy.zapisane == 7
    assert drugi.zapisane == 0
    assert drugi.pominiete_jako_powtorzenia == 7
    assert len(magazyn.wczytaj_wpisy(p)) == 7


def test_dopisany_wiersz_wchodzi_przy_ponownym_imporcie(tmp_path, p):
    """Deduplikacja liczy KROTNOSC, nie unikalnosc: dolozenie trzeciej takiej
    samej kawy ma dopisac jedna, a nie zero i nie trzy."""
    arkusz = _zbuduj_arkusz(tmp_path / "przed.xlsx")
    import_arkusza.importuj(p, arkusz)
    wiecej = list(WPISY_SYNTETYCZNE) + [
        (datetime.datetime(2025, 10, 3), "Wydatki", "Jedzenie poza domem", 12.00)]
    arkusz2 = _zbuduj_arkusz(tmp_path / "po.xlsx", wiecej)
    drugi = import_arkusza.importuj(p, arkusz2)
    assert drugi.zapisane == 1
    assert len(magazyn.wczytaj_wpisy(p)) == 8


# --- slownik kategorii ----------------------------------------------------

def test_import_dodaje_kategorie_ktorych_nie_bylo_w_slowniku(tmp_path, p):
    przed = len(magazyn.wczytaj_kategorie(p))
    arkusz = _zbuduj_arkusz(tmp_path / "syntetyk.xlsx")
    import_arkusza.importuj(p, arkusz)
    po = len(magazyn.wczytaj_kategorie(p))
    assert przed == sum(
        len(v) for v in budzet.KATEGORIE_STARTOWE.values())
    assert po >= przed


def test_import_pomija_sloty_placeholdery_spoza_tabel(arkusz, p):
    """W prawdziwym pliku poza zakresami tabel lezy osiem slotow, w tym siedem
    powtorzen "Inne". Wciagniecie ich dalo by duplikaty i nazwy-zachety."""
    import_arkusza.importuj(p, arkusz)
    nazwy = set(magazyn.wczytaj_kategorie(p)["nazwa"])
    assert not any("Wprowadź" in n for n in nazwy)


def test_kategorie_z_arkusza_czytane_z_zakresow_tabel(arkusz):
    kategorie = import_arkusza.wczytaj_kategorie_z_arkusza(arkusz)
    assert set(kategorie) == set(budzet.TYPY)
    assert kategorie["Wydatki"] == ["Transport", "Zakupy",
                                    "Jedzenie poza domem", "Kredyty", "Inne"]
    assert kategorie["Przychody"] == ["Wynagrodzenie", "Inne"]


# --- plan budzetu ---------------------------------------------------------

def test_import_wciaga_plan_per_miesiac_i_kategoria(arkusz, p):
    import_arkusza.importuj(p, arkusz)
    plan = magazyn.wczytaj_plan(p, rok=2025, miesiac=10)
    pozycje = {(w["typ"], w["kategoria"]): w["kwota_planowana"]
               for _, w in plan.iterrows()}
    assert pozycje[("Wydatki", "Transport")] == Decimal("100")
    assert pozycje[("Wydatki", "Zakupy")] == Decimal("101")
    assert pozycje[("Przychody", "Wynagrodzenie")] == Decimal("100")
    listopad = magazyn.wczytaj_plan(p, rok=2025, miesiac=11)
    pozycje_11 = {(w["typ"], w["kategoria"]): w["kwota_planowana"]
                  for _, w in listopad.iterrows()}
    assert pozycje_11[("Wydatki", "Transport")] == Decimal("200")


def test_ponowny_import_planu_nadpisuje_nie_duplikuje(arkusz, p):
    import_arkusza.importuj(p, arkusz)
    ile = len(magazyn.wczytaj_plan(p))
    import_arkusza.importuj(p, arkusz)
    assert len(magazyn.wczytaj_plan(p)) == ile


def test_plan_pomija_sloty_placeholdery(arkusz, p):
    import_arkusza.importuj(p, arkusz)
    plan = magazyn.wczytaj_plan(p)
    assert not any("Wprowadź" in k for k in plan["kategoria"])
    assert Decimal("999") not in set(plan["kwota_planowana"])


# --- kontrola dwustronna na prawdziwym pliku ------------------------------

@pytest.mark.arkusz
def test_prawdziwy_arkusz_ma_204_wpisy_i_23_kategorie():
    """Kontrola dwustronna poza bramka domyslna: plik lezy w iCloud, poza repo.

    Bez niej "zaimportowano 0 wpisow" przeszlo by jako sukces.
    """
    assert SCIEZKA_PRAWDZIWEGO.exists(), f"brak pliku: {SCIEZKA_PRAWDZIWEGO}"
    wpisy, powody = import_arkusza.wczytaj_wpisy_z_arkusza(SCIEZKA_PRAWDZIWEGO)
    assert len(wpisy) == 204, f"przeczytano {len(wpisy)}, oczekiwano 204"
    assert not powody, f"nieoczekiwane odrzucenia: {powody}"

    kategorie = import_arkusza.wczytaj_kategorie_z_arkusza(SCIEZKA_PRAWDZIWEGO)
    assert sum(len(v) for v in kategorie.values()) == 23
    assert len(kategorie["Przychody"]) == 4
    assert len(kategorie["Wydatki"]) == 15
    assert len(kategorie["Oszczędności"]) == 4

    typy = {}
    for wpis in wpisy:
        typy[wpis.typ] = typy.get(wpis.typ, 0) + 1
    assert typy == {"Wydatki": 181, "Oszczędności": 12, "Przychody": 11}
